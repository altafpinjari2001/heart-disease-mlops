"""Generate reports/MLOps_Assignment1_Report.docx from the pipeline outputs.

    pip install python-docx
    python scripts/build_report.py

Screenshots that are not yet present in screenshots/ are rendered as a
labelled placeholder box, so re-running the script after capturing them
produces the final report.
"""

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "reports" / "figures"
SHOT = ROOT / "screenshots"
OUT = ROOT / "reports" / "MLOps_Assignment1_Report.docx"

REPO_URL = "https://github.com/altafpinjari2001/heart-disease-mlops"
VIDEO_URL = ("https://drive.google.com/file/d/1N9qXK5QjRM7D30F2cs8cYlieP3KHAZhm/"
             "view?usp=sharing")
STUDENT = "Altaf Anwar Pinjari  |  2025AE05638"

ACCENT = RGBColor(0x1F, 0x6F, 0x8B)
meta = json.loads((ROOT / "models" / "model_metadata.json").read_text())
comparison = pd.read_csv(ROOT / "reports" / "model_comparison.csv",
                         index_col=0)
runs = pd.read_csv(ROOT / "reports" / "experiment_runs.csv")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(1.9))
sec.top_margin = sec.bottom_margin = Cm(1.6)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(3)
for lvl, size in ((1, 15), (2, 12)):
    st = styles[f"Heading {lvl}"]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = ACCENT
    st.font.bold = True


def add_link(paragraph, url, text, size=None, bold=False):
    """Append a clickable, underlined hyperlink run to a paragraph."""
    rid = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
             "hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:color", {"w:val": "1F6F8B"}), ("w:u", {"w:val": "single"})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


def para(text, bold_lead=None, size=None, align=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    r = p.add_run(text.replace("`", ""))
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1].replace("`", ""))
        else:
            p.add_run(it.replace("`", ""))


def code(text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F1F4F6")
        p._p.get_or_add_pPr().append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x66, 0x70)


def image(path, width_cm, cap):
    path = Path(path)
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        t = doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        c.width = Cm(width_cm)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.paragraphs[0].add_run(
            f"\n[ Insert screenshot: screenshots/{path.name} ]\n")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    caption(cap)


def image_pair(left, right, width_cm, cap):
    """Two figures side by side in a borderless 2-column table."""
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, path in zip(t.rows[0].cells, (Path(left), Path(right))):
        cell.width = Cm(width_cm + 0.2)
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.exists():
            par.add_run().add_picture(str(path), width=Cm(width_cm))
        else:
            r = par.add_run(f"\n[ Insert screenshot:\nscreenshots/{path.name} ]\n")
            r.bold = True
            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    caption(cap)


def table(header, rows, widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ title
para("Machine Learning Operations (AIMLCZG523)  |  Assignment 01", size=11,
     align=WD_ALIGN_PARAGRAPH.CENTER)
t = para("Heart Disease Risk Prediction: an End-to-End MLOps Pipeline",
         size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
t.runs[0].bold = True
t.runs[0].font.color.rgb = ACCENT
para(STUDENT, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
rp = para("Code repository: ", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
rp.paragraph_format.space_after = Pt(0)
add_link(rp, REPO_URL, REPO_URL, size=10)
vp = para("Demo video: ", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
add_link(vp, VIDEO_URL, VIDEO_URL, size=10)
rule = vp._p.get_or_add_pPr()
bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"),
             ("w:color", "1F6F8B")):
    bottom.set(qn(k), v)
bdr.append(bottom)
rule.append(bdr)

# ------------------------------------------------------------ 1 overview
doc.add_heading("1. Problem, objective and summary", level=1)
para("The goal is to predict whether a patient has heart disease from 13 "
     "routinely collected clinical attributes and to expose the model as a "
     "reproducible, containerised, monitored REST API. The dataset is the "
     "Cleveland subset of the UCI Heart Disease database (303 patients). The "
     "original target `num` (0-4) is binarised: 0 = no disease, 1-4 = "
     "disease present.")
m = meta["metrics"]
table(["Outcome", "Value"], [
    ["Champion model", "Logistic regression (L2, C = 0.3, "
                       "class_weight = balanced)"],
    ["Selection metric (5-fold CV ROC-AUC, train split)",
     f"{m['cv_roc_auc']:.3f}"],
    ["Held-out test ROC-AUC / accuracy",
     f"{m['test_roc_auc']:.3f} / {m['test_accuracy']:.3f}"],
    ["Held-out test recall / precision",
     f"{m['test_recall']:.3f} / {m['test_precision']:.3f}"],
    ["Runs tracked in MLflow", f"{len(runs) + 4} (64 nested tuning runs)"],
    ["Unit tests", "28 pytest tests, flake8 clean"],
    ["Serving", "FastAPI in Docker, 2-5 replicas on Kubernetes behind a "
                "LoadBalancer + Ingress, Prometheus/Grafana monitoring"],
], widths=[7, 9.6])

doc.add_heading("1.1 Architecture", level=2)
image(FIG / "architecture.png", 14,
      "Figure 1 - System architecture: data & experimentation, CI/CD, "
      "serving on Kubernetes and monitoring.")
para("Design principle: every stage is a plain Python entry point "
     "(`scripts/download_data.py`, `python -m heart_risk.train`, "
     "`uvicorn api.main:app`) so the same commands run on a laptop, in "
     "GitHub Actions and inside the container. All preprocessing lives inside "
     "a single scikit-learn Pipeline object, which is the artefact that is "
     "versioned, containerised and served - this removes train/serve skew.")

# ------------------------------------------------------------- 2 setup
doc.add_heading("2. Setup and installation", level=1)
para("Requirements: Python 3.11, Git; Docker Desktop with Kubernetes (or "
     "Minikube) for the deployment steps. All library versions are "
     "pinned in requirements.txt; requirements-api.txt is the minimal "
     "runtime subset used by the container (same pins, so the pickled "
     "pipeline always loads).")
code("""
git clone https://github.com/altafpinjari2001/heart-disease-mlops.git && cd heart-disease-mlops
python -m venv .venv && .venv\\Scripts\\activate      # Linux: source .venv/bin/activate
pip install -r requirements-dev.txt
python scripts/download_data.py      # download + clean the UCI data
python -m heart_risk.eda             # EDA figures
python -m heart_risk.train           # tune 3 models, log to MLflow, save champion
pytest -v && flake8                  # tests + lint
uvicorn api.main:app --port 8000     # serve locally -> http://127.0.0.1:8000/docs
""")
table(["Folder / file", "Purpose"], [
    ["heart_risk/", "Package: config, data cleaning, features, training, "
                    "evaluation, inference, EDA"],
    ["api/", "FastAPI app, request/response schemas, logging + metrics, "
             "/monitor page"],
    ["notebooks/", "01_eda, 02_model_training, 03_inference (executed)"],
    ["data/raw, data/processed", "Raw UCI file and cleaned CSV"],
    ["models/", "heart_model.joblib, MLflow model, model_metadata.json"],
    ["mlruns/", "MLflow tracking store (all runs, metrics, artefacts)"],
    ["test/", "pytest unit tests (data, features, model, API)"],
    [".github/workflows/ci.yml", "Lint -> test -> train -> Docker pipeline"],
    ["Dockerfile, docker-compose.yml", "Serving image; local API + "
                                       "Prometheus + Grafana stack"],
    ["k8s/", "Namespace, Deployment, Service, Ingress, HPA, monitoring"],
    ["monitoring/", "Prometheus config, alert rules, Grafana dashboard"],
    ["screenshots/, reports/", "Evidence and this report"],
], widths=[5.2, 11.4])

# --------------------------------------------------------------- 3 EDA
doc.add_heading("3. Data acquisition, cleaning and EDA", level=1)
doc.add_heading("3.1 Cleaning", level=2)
bullets([
    ("Missing values: ", "the raw file marks missing values with '?'. Only "
     "`ca` (4 rows) and `thal` (2 rows) are affected. They are kept as NaN "
     "in the cleaned CSV and imputed inside the pipeline (median for "
     "numeric, most-frequent for categorical). Imputing in the pipeline "
     "means the statistics are learned on training folds only and the "
     "identical rule is applied to API requests - the API accepts requests "
     "with `ca`/`thal` omitted."),
    ("Validity rules: ", "values outside physiological ranges (e.g. chol "
     "< 100) or invalid category codes are converted to NaN; none were found "
     "in this subset but the rule protects future data."),
    ("Types and duplicates: ", "all columns cast to numeric, binary columns "
     "to int64 (fixed width so the MLflow signature is platform "
     "independent), duplicate rows dropped (none)."),
])
doc.add_heading("3.2 Findings", level=2)
image_pair(FIG / "eda" / "01_class_balance.png",
           FIG / "eda" / "02_missing_values.png", 7.2,
           "Figure 2 - Class balance (164 healthy vs 139 diseased, 54 / 46 %) "
           "and missing values in the raw file.")
para("The classes are only mildly imbalanced, so stratified splits are "
     "sufficient; class_weight='balanced' is still explored during tuning.")
image(FIG / "eda" / "03_histograms.png", 12,
      "Figure 3 - Continuous features by diagnosis.")
bullets([
    "Diseased patients reach a markedly lower maximum heart rate (thalach) "
    "and show higher exercise-induced ST depression (oldpeak).",
    "Cholesterol and resting blood pressure overlap strongly between "
    "classes - weak individual predictors.",
    "The number of coloured vessels (ca) shows a monotonic risk increase, "
    "so it is treated as an ordinal numeric feature.",
])
image_pair(FIG / "eda" / "04_correlation_heatmap.png",
           FIG / "eda" / "07_hr_vs_age.png", 8.2,
           "Figure 4 - Spearman correlation heatmap (left) and max heart rate "
           "vs age (right): most diseased patients fall below 85 % of the "
           "age-predicted maximum.")
image(FIG / "eda" / "05_categorical_disease_rate.png", 13,
      "Figure 5 - Disease rate per category (dashed line = overall rate).")
para("Chest pain type 4 (asymptomatic), reversible thalassemia defect, flat "
     "ST slope and exercise-induced angina all have disease rates above "
     "65 %. Because these are nominal codes (cp = 4 is not 'four times' "
     "cp = 1), cp, restecg, slope and thal are one-hot encoded.")
para("Figure 4 (right) motivates the engineered feature hr_reserve_ratio = "
     "thalach / (220 - age): the fraction of the age-predicted maximum heart "
     "rate reached during the stress test. It separates the age effect from "
     "the cardiac response. Box plots (reports/figures/eda/06_boxplots.png) "
     "show a few high-cholesterol outliers that are physiologically "
     "plausible and were therefore kept.")

# ---------------------------------------------------------- 4 modelling
doc.add_heading("4. Feature engineering and model development", level=1)
doc.add_heading("4.1 Preprocessing pipeline", level=2)
table(["Step", "Columns", "Transformation"], [
    ["ClinicalFeatureBuilder", "age, thalach", "adds hr_reserve_ratio"],
    ["numeric", "age, trestbps, chol, thalach, oldpeak, ca, "
                "hr_reserve_ratio", "median impute -> StandardScaler"],
    ["binary", "sex, fbs, exang", "most-frequent impute (kept 0/1)"],
    ["categorical", "cp, restecg, slope, thal",
     "most-frequent impute -> one-hot with fixed UCI category levels, "
     "handle_unknown='ignore'"],
], widths=[3.6, 6.2, 6.8])
para("13 raw inputs become 23 model features. Fixing the category levels "
     "from the UCI documentation guarantees a stable feature layout even if "
     "a level is missing from a training fold or a request.")
doc.add_heading("4.2 Models, tuning and selection", level=2)
bullets([
    ("Split: ", "80 / 20 stratified train/test split (242 / 61 patients, "
     "random_state = 42). The test split is never used for any decision."),
    ("Cross-validation: ", "StratifiedKFold(5, shuffle) on the training "
     "split; metrics accuracy, precision, recall, F1 and ROC-AUC."),
    ("Logistic regression: ", "C in {0.01 ... 3}, penalty L1/L2, "
     "class_weight None/balanced (24 configurations)."),
    ("Random forest: ", "300 trees, max_depth {3, 5, 8, None}, "
     "min_samples_leaf {1, 4, 8}, max_features {sqrt, 0.5} (24)."),
    ("Gradient boosting: ", "n_estimators {100, 200}, learning_rate "
     "{0.03, 0.1}, max_depth {2, 3}, subsample {0.7, 1.0} (16)."),
    ("Selection rule: ", "GridSearchCV refits on ROC-AUC (a threshold-free "
     "ranking metric suited to a risk score); the champion is the family with "
     "the highest mean CV ROC-AUC."),
])
rows = []
for fam, r in comparison.iterrows():
    rows.append([fam.replace("_", " "), f"{r.cv_roc_auc:.3f}",
                 f"{r.cv_accuracy:.3f}", f"{r.cv_recall:.3f}",
                 f"{r.test_roc_auc:.3f}", f"{r.test_accuracy:.3f}",
                 f"{r.test_precision:.3f}", f"{r.test_recall:.3f}"])
table(["Model", "CV AUC", "CV acc", "CV recall", "Test AUC", "Test acc",
       "Test prec", "Test recall"], rows,
      widths=[3.6, 1.7, 1.7, 1.8, 1.8, 1.8, 1.9, 1.9], font=8.5)
image(FIG / "model_comparison.png", 11,
      "Figure 6 - Cross-validated metrics of the tuned model families.")
para("Logistic regression wins on CV ROC-AUC (0.907) and accuracy. The tree "
     "ensembles score slightly higher on the 61-patient test split, but that "
     "difference (+/- 1-2 patients) is within noise and choosing on it would "
     "overfit the test set. With 242 training rows and a largely additive "
     "signal, the simpler, interpretable linear model is the better "
     "production choice: it is fast, small, well calibrated and its "
     "coefficients can be explained to clinicians.")
image_pair(FIG / "logistic_regression" / "confusion_matrix.png",
           FIG / "logistic_regression" / "feature_importance.png", 6.8,
           "Figure 7 - Champion confusion matrix on the test split (left) and "
           "coefficients on standardised features (right).")
para("Recall of 0.93 on the test split (26 of 28 diseased patients found) is "
     "the clinically important number: a screening model should prefer false "
     "alarms over missed cases. The most influential features - ca, sex, "
     "thal, cp = 4 and exang - match the EDA. thalach and hr_reserve_ratio "
     "are correlated and share their weight.")

# --------------------------------------------------- 5 experiment tracking
doc.add_heading("5. Experiment tracking (MLflow)", level=1)
bullets([
    ("Parent run per model family: ", "best hyper-parameters, CV metrics, "
     "test metrics, tags (model_family, stage, package_version), plots "
     "(confusion matrix, ROC, PR, feature importance), the full CV results "
     "CSV and the logged sklearn model with input signature and example."),
    ("Nested child run per grid point: ", "64 runs with parameters and mean "
     "/ std of all five CV metrics - every configuration tried is "
     "traceable."),
    ("champion-selection run: ", "tags the winning family and run id and "
     "logs the comparison chart, comparison CSV and model_metadata.json."),
])
image_pair(SHOT / "03_mlflow_experiment_runs.png",
           SHOT / "04_mlflow_lr_run_details.png", 8.2,
           "Figure 8 - MLflow experiment 'heart-disease-risk' (left) and the "
           "champion run with parameters and metrics (right).")
image(FIG / "experiment_tracking_summary.png", 13,
      "Figure 9 - All 64 tuning runs read back from MLflow "
      "(scripts/experiment_summary.py).")
para("The tracking data also explains the tuning behaviour: strong L1 "
     "regularisation (C <= 0.03) removes every coefficient (AUC 0.5), and "
     "logistic regression plateaus at ~0.90-0.91 from C = 0.1. Random-forest "
     "configurations are tightly clustered (0.87-0.90), while gradient "
     "boosting with deeper trees overfits.")

# ------------------------------------------------- 6 packaging
doc.add_heading("6. Model packaging and reproducibility", level=1)
bullets([
    ("models/heart_model.joblib: ", "the complete Pipeline (feature builder "
     "+ imputers + scaler + encoder + classifier). Raw JSON fields in, "
     "probability out - used by the API."),
    ("models/mlflow_model/: ", "the same pipeline in MLflow format "
     "(MLmodel, conda.yaml, python_env.yaml, requirements.txt, signature, "
     "input example) - loadable with mlflow.pyfunc.load_model."),
    ("models/model_metadata.json: ", "version, family, MLflow run id, best "
     "params, metrics, feature list, decision threshold, training time."),
    ("Determinism: ", "fixed random_state for split, CV and models; "
     "re-training from scratch reproduced identical metrics."),
    ("Environment: ", "pinned requirements.txt (+ requirements-api.txt for "
     "serving); Python 3.11 locally, in CI and in the image."),
])

# ----------------------------------------------------- 7 CI/CD + tests
doc.add_heading("7. CI/CD pipeline and automated testing", level=1)
table(["Test module", "What is verified"], [
    ["test_data.py (7)", "binary target mapping, '?' parsing, missing "
                         "values kept for imputation, out-of-range and "
                         "invalid codes -> NaN, de-duplication, download "
                         "skipped when cached"],
    ["test_features.py (5)", "hr_reserve_ratio value, no NaNs after "
                             "preprocessing, stable 23-column output even "
                             "with missing categories, scaling, feature "
                             "names"],
    ["test_model.py (5)", "models learn signal (AUC > 0.75), metric "
                          "values, prediction contract, save/load "
                          "determinism, plots written"],
    ["test_api.py (11)", "/health, /predict output + request id header, "
                         "optional fields, 422 on invalid / unknown fields, "
                         "batch, /metrics, /monitor, /model-info"],
], widths=[3.8, 12.8])
image(SHOT / "06_local_tests_lint_coverage.png", 8.5,
      "Figure 10 -flake8 (0 issues) and 28 passing tests with coverage.")
para("GitHub Actions workflow (.github/workflows/ci.yml), triggered on push, "
     "pull request and manually:")
table(["Job", "Steps", "Artefacts"], [
    ["lint", "flake8 over all code", "flake8-report"],
    ["test (needs lint)", "pytest -v, JUnit XML, coverage", "test-reports"],
    ["train (needs test)", "download data, full training with MLflow, "
                           "quality gate: test ROC-AUC >= 0.85",
     "trained-model, mlflow-runs, training.log"],
    ["docker (needs train)", "build image with the freshly trained model, "
                             "run it, wait for /health, POST "
                             "sample_request.json, assert response",
     "docker-smoke-test, container.log, docker-image"],
], widths=[3.4, 8.4, 4.8])
para("Jobs are chained with needs:, so any lint error, failing test, "
     "training crash, quality-gate miss or broken container stops the "
     "pipeline, and each failing step keeps its full log in the run.")
image(SHOT / "07_github_actions_pipeline.png", 10,
      "Figure 11 - GitHub Actions run #6: lint, tests, training and Docker jobs all green, 6 artefacts, no warnings.")
image(SHOT / "07b_github_actions_docker_job.png", 14,
      "Figure 12 - CI Docker job log: image build, container start, "
      "/health and /predict smoke test, structured container logs.")

# ------------------------------------------------ 8 container + API
doc.add_heading("8. Model containerisation and API", level=1)
bullets([
    ("Image: ", "multi-stage build on python:3.11-slim (wheels built in a "
     "builder stage), runs as non-root uid 1001, HEALTHCHECK on /health, "
     "only heart_risk/, api/ and the model files are copied."),
    ("POST /predict: ", "validated JSON (Pydantic: ranges, allowed codes, "
     "unknown fields rejected) -> prediction, label, confidence (probability "
     "of the predicted class), risk_probability, model_version, request_id."),
    ("Other endpoints: ", "/predict/batch (up to 500 patients), /health, "
     "/model-info, /metrics (Prometheus), /monitor (live dashboard), /docs."),
])
para("The image was built and run locally with Docker Desktop 29.9 "
     "(final image 258 MB) and is also built and smoke-tested in every CI "
     "run (Figure 12).")
image(SHOT / "14_docker_build.png", 8,
      "Figure 13 - Local multi-stage docker build of heart-risk-api:1.0.0.")
image(SHOT / "15_docker_run_predict.png", 13.5,
      "Figure 14 - Container running (docker ps), /health and /predict "
      "answered from the container, JSON logs via docker logs.")

# ------------------------------------------------ 9 deployment
doc.add_heading("9. Production deployment (Kubernetes)", level=1)
table(["Manifest", "Key settings"], [
    ["00-namespace.yaml", "isolated namespace heart-risk"],
    ["01-deployment.yaml", "2 replicas, RollingUpdate (maxUnavailable 0), "
                           "readiness + liveness probes on /health, CPU / "
                           "memory requests and limits, non-root, Prometheus "
                           "scrape annotations"],
    ["02-service.yaml", "type LoadBalancer, port 8080 -> 8000"],
    ["03-ingress.yaml", "NGINX Ingress, host heart-risk.local"],
    ["04-hpa.yaml", "autoscale 2-5 replicas at 70 % CPU (metrics-server)"],
    ["monitoring/*.yaml", "Prometheus (RBAC + pod discovery) and Grafana "
                          "with provisioned datasource and dashboard"],
], widths=[4.2, 12.4])
para("The stack was deployed to Docker Desktop's built-in Kubernetes "
     "(single node, v1.36.1) with one script that builds the image, installs "
     "the NGINX ingress controller and metrics-server, and applies every "
     "manifest plus the monitoring stack:")
code("""
bash scripts/deploy_k8s.sh        # Windows: scripts\\deploy_k8s.ps1 (also supports Minikube)
kubectl -n heart-risk get deploy,pods,svc,ingress,hpa
curl -X POST http://localhost:8080/predict -H "Content-Type: application/json" -d @sample_request.json
kubectl -n ingress-nginx port-forward svc/ingress-nginx-controller 8088:80
curl -H "Host: heart-risk.local" http://localhost:8088/predict ...
""")
para("Access instructions (local cluster, no public URL): the LoadBalancer "
     "publishes the API on http://localhost:8080. Docker Desktop's load "
     "balancer serves a single LoadBalancer service, so the ingress "
     "controller runs as NodePort and is reached through a port-forward on "
     "localhost:8088 with the Host header heart-risk.local; a request without "
     "the host returns 404, proving host-based routing. On Minikube the same "
     "script works with minikube tunnel. Under the load test the HPA scaled "
     "the API from 2 to 4 replicas, and Prometheus discovered every new pod.")
image(SHOT / "16_k8s_resources.png", 12,
      "Figure 15 - Node, Deployments, Pods, Services, Ingress and HPA "
      "(scaled 2 -> 4 under load) in namespace heart-risk.")
image(SHOT / "17_k8s_predict_loadbalancer.png", 13,
      "Figure 16 - Predictions served through the LoadBalancer and the "
      "NGINX Ingress, plus pod logs.")

# ------------------------------------------------ 10 monitoring
doc.add_heading("10. Monitoring and logging", level=1)
bullets([
    ("Request logging: ", "middleware assigns a request id (returned in "
     "x-request-id), times each request and writes one JSON line per request "
     "and per prediction (inputs, label, probability) to stdout and "
     "/app/logs/api.log - ready for kubectl logs or a log shipper."),
    ("Prometheus metrics: ", "heart_api_requests_total{endpoint,status}, "
     "heart_api_request_latency_seconds (histogram), "
     "heart_api_predictions_total{label}, heart_api_risk_probability "
     "(histogram), heart_api_inference_errors_total, heart_api_model_info."),
    ("Alert rules: ", "API down, 5xx rate > 5 %, p95 /predict latency > "
     "500 ms, and a simple drift signal - positive-prediction share above "
     "80 % for 15 min (training prevalence is 46 %)."),
    ("Dashboards: ", "Grafana dashboard (status, totals, positive share, "
     "error rate, latency percentiles, request rate, predictions per label, "
     "risk distribution) and a dependency-free /monitor page built into the "
     "API (screenshot in screenshots/11_api_monitor_dashboard.png)."),
])
image(SHOT / "18_grafana_dashboard.png", 14,
      "Figure 17 - Grafana dashboard in the cluster after ~3,200 predictions "
      "through the LoadBalancer: 0 % 5xx, p95 latency ~50 ms, positive share "
      "46 % (matches training prevalence). The early latency spike is the "
      "load burst before the HPA added replicas.")
image(SHOT / "19_prometheus_targets.png", 10,
      "Figure 18 - Prometheus pod discovery: all 4 API replicas UP.")

# ------------------------------------------------ 11 conclusion
doc.add_heading("11. Conclusion and future work", level=1)
para("The project delivers a reproducible path from raw UCI data to a "
     "monitored API: one pipeline object carries all preprocessing, every "
     "experiment is tracked, CI re-trains and gates the model before building "
     "the image, and Kubernetes runs it with health checks, autoscaling and "
     "metrics. Next steps: a model registry with staged promotion, pushing "
     "the image to a container registry from CI, decision-threshold tuning "
     "with clinicians, and statistical input-drift tests on logged requests.")
add_link(para("Code repository: "), REPO_URL, REPO_URL, bold=True)
add_link(para("Demo video (pipeline walkthrough): "), VIDEO_URL, VIDEO_URL, bold=True)

doc.save(OUT)
print(OUT)
